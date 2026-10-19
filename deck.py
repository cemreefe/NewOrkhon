"""Build a vector presentation from actual, HarfBuzz-shaped New Orkhon outlines."""

from __future__ import annotations

import argparse
import html
import json
import sys
import unicodedata
from dataclasses import asdict, dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

import fitz
import uharfbuzz as hb
from fontTools.pens.basePen import BasePen
from fontTools.pens.boundsPen import BoundsPen
from fontTools.ttLib import TTFont
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT / "typeface" if (ROOT / "typeface").exists() else ROOT.parent / "new-orkhon"
ASSETS = ROOT / "assets"
OUT = ROOT / "output"
PREVIEWS = ROOT / "previews"
W, H, TOTAL = 1600, 900, 14
PAPER, INK = "#FFFFFF", "#000000"
MUTED = RULE = WARM_GREY = WINE = INK
REGULAR = ASSETS / "InstrumentSans-Regular.ttf"
MEDIUM = ASSETS / "InstrumentSans-Medium.ttf"
MONO = ASSETS / "JetBrainsMono-Regular.ttf"
if not MONO.exists():
    MONO = Path("/usr/share/fonts/truetype/jetbrains-mono/JetBrainsMono-Regular.ttf")
UPPER = PROJECT / "NewOrkhonText-Regular.ttf"
LOWER = PROJECT / "NewOrkhonTextLowercase-Regular.ttf"
SANS_UPPER = PROJECT / "NewOrkhonSans-Regular.ttf"
SANS_LOWER = PROJECT / "NewOrkhonSansLowercase-Regular.ttf"
REFERENCE = PROJECT / "NotoSansOldTurkic-Regular.ttf"
SCALE = Inches(13 + 1 / 3) / W
sys.path.insert(0, str(PROJECT))
from build import LETTERS, VERSION  # noqa: E402
from showcase import historical  # noqa: E402


class OutlinePen(BasePen):
    """Keep exact cubic outlines; convert quadratic segments once."""

    def __init__(self, glyph_set: Any) -> None:
        super().__init__(glyph_set)
        self.commands: list[tuple[Any, ...]] = []

    def _moveTo(self, p: tuple[float, float]) -> None:
        self.commands.append(("M", *p))

    def _lineTo(self, p: tuple[float, float]) -> None:
        self.commands.append(("L", *p))

    def _curveToOne(self, p1: Any, p2: Any, p3: Any) -> None:
        self.commands.append(("C", *p1, *p2, *p3))

    def _qCurveToOne(self, p1: Any, p2: Any) -> None:
        p0 = self._getCurrentPoint()
        assert p0 is not None
        self._curveToOne(
            (p0[0] + 2 / 3 * (p1[0] - p0[0]), p0[1] + 2 / 3 * (p1[1] - p0[1])),
            (p2[0] + 2 / 3 * (p1[0] - p2[0]), p2[1] + 2 / 3 * (p1[1] - p2[1])),
            p2,
        )

    def _closePath(self) -> None:
        self.commands.append(("Z",))

    def _endPath(self) -> None:
        pass


@lru_cache
def font_data(path: str) -> tuple[TTFont, hb.Font]:
    tt = TTFont(path)
    font = hb.Font(hb.Face(Path(path).read_bytes()))
    upm = tt["head"].unitsPerEm
    font.scale = (upm, upm)
    return tt, font


@lru_cache
def glyph_outline(path: str, name: str) -> tuple[list[Any], tuple[float, ...]]:
    tt, _ = font_data(path)
    glyphs = tt.getGlyphSet()
    pen = OutlinePen(glyphs)
    glyphs[name].draw(pen)
    bounds = BoundsPen(glyphs)
    glyphs[name].draw(bounds)
    return pen.commands, bounds.bounds or (0, 0, 0, 0)


def transform(commands: list[Any], x: float, y: float, sx: float, sy: float) -> list[Any]:
    result = []
    for command in commands:
        values: list[Any] = [command[0]]
        for i in range(1, len(command), 2):
            values.extend((x + command[i] * sx, y + command[i + 1] * sy))
        result.append(tuple(values))
    return result


def shape(
    path: Path, text: str, features: dict[str, Any] | None = None, rtl: bool = True
) -> tuple[list[Any], float, tuple[float, ...], list[str]]:
    tt, font = font_data(str(path))
    buffer = hb.Buffer()
    buffer.add_str(text)
    if rtl:
        buffer.direction, buffer.script, buffer.language = "rtl", "orkh", "otk"
    else:
        buffer.guess_segment_properties()
    hb.shape(font, buffer, features or {})
    commands, names = [], []
    x, y = 0.0, 0.0
    x0, y0, x1, y1 = float("inf"), float("inf"), -float("inf"), -float("inf")
    for info, position in zip(buffer.glyph_infos, buffer.glyph_positions, strict=True):
        name = tt.getGlyphName(info.codepoint)
        assert name != ".notdef", (path, text)
        names.append(name)
        glyph, bounds = glyph_outline(str(path), name)
        gx, gy = x + position.x_offset, y + position.y_offset
        commands.extend(transform(glyph, gx, gy, 1, 1))
        if glyph:
            x0, y0 = min(x0, gx + bounds[0]), min(y0, gy + bounds[1])
            x1, y1 = max(x1, gx + bounds[2]), max(y1, gy + bounds[3])
        x += position.x_advance
        y += position.y_advance
    return commands, x, (x0, y0, x1, y1), names


@dataclass
class Slide:
    title: str
    bg: str = PAPER
    elements: list[dict[str, Any]] = field(default_factory=list)
    notes: str = ""

    def rect(self, x: float, y: float, w: float, h: float, color: str) -> None:
        self.elements.append(dict(kind="rect", x=x, y=y, w=w, h=h, color=color))

    def line(self, x1: float, y1: float, x2: float, y2: float, color: str = RULE, width: float = 1) -> None:
        self.elements.append(dict(kind="line", x=x1, y=y1, x2=x2, y2=y2, color=color, width=width))

    def path(
        self, commands: list[Any], color: str | None, name: str = "Vector artwork",
        stroke: str | None = None, width: float = 1.3,
    ) -> None:
        self.elements.append(dict(
            kind="path", commands=commands, color=color, name=name, stroke=stroke, width=width,
        ))

    def text(
        self, text: str, x: float, baseline: float, size: float = 26,
        color: str = INK, font: Path = REGULAR, align: str = "left",
    ) -> float:
        commands, advance, bounds, _ = shape(font, text, {"kern": 0, "liga": 0}, rtl=False)
        factor = size / font_data(str(font))[0]["head"].unitsPerEm
        width = advance * factor
        left = x - (width / 2 if align == "center" else width if align == "right" else 0)
        self.elements.append(dict(
            kind="text", text=text, x=left, y=baseline, size=size, color=color,
            font=str(font), w=width, commands=transform(commands, left, baseline, factor, -factor),
            bounds=[left + bounds[0] * factor, baseline - bounds[3] * factor,
                    left + bounds[2] * factor, baseline - bounds[1] * factor],
        ))
        return width

    def run(
        self, text: str, x: float, baseline: float, size: float,
        color: str = INK, path: Path = LOWER,
        features: dict[str, Any] | None = None, align: str = "left",
        name: str | None = None, outline: bool = False, width: float = 1.3,
    ) -> tuple[float, tuple[float, ...]]:
        commands, advance, bounds, names = shape(path, text, features)
        factor = size / font_data(str(path))[0]["head"].unitsPerEm
        advance_width = advance * factor
        left = x - (
            advance_width / 2 if align == "center" else advance_width if align == "right" else 0
        )
        self.path(
            transform(commands, left, baseline, factor, -factor),
            None if outline else color, name or text, color if outline else None, width,
        )
        self.elements[-1].update(source=text, font=str(path), features=features or {}, glyphs=names)
        return advance_width, (
            left + bounds[0] * factor, baseline - bounds[3] * factor,
            left + bounds[2] * factor, baseline - bounds[1] * factor,
        )

    def furniture(self, number: int, section: str, dark: bool = False) -> None:
        color = PAPER if dark else INK
        self.text("NEW ORKHON", 88, 60, 16, color, MEDIUM)
        self.text(section.upper(), 1512, 60, 14, color, MONO, "right")
        self.text("YENİ ORHUN", 88, 850, 13, color, MONO)
        self.text(f"{number:02d} / {TOTAL:02d}", 1512, 850, 13, color, MONO, "right")


def draw_path_pdf(
    c: canvas.Canvas, commands: list[Any], color: str | None,
    stroke: str | None = None, width: float = 1.3,
) -> None:
    p = c.beginPath()
    for row in commands:
        op, v = row[0], row[1:]
        if op == "M":
            p.moveTo(v[0], v[1])
        elif op == "L":
            p.lineTo(v[0], v[1])
        elif op == "C":
            p.curveTo(*v)
        elif op == "Z":
            p.close()
    if color:
        c.setFillColor(color)
    if stroke:
        c.setStrokeColor(stroke)
        c.setLineWidth(width)
    c.drawPath(p, stroke=int(stroke is not None), fill=int(color is not None), fillMode=1)


def export_pdf(slides: list[Slide], path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=(960, 540), pageCompression=1, invariant=1)
    c.setTitle("New Orkhon — Typeface Presentation")
    c.setAuthor("New Orkhon")
    for slide in slides:
        c.saveState()
        c.translate(0, 540)
        c.scale(.6, -.6)
        c.setFillColor(slide.bg)
        c.rect(0, 0, W, H, stroke=0, fill=1)
        for el in slide.elements:
            if el["kind"] in ("path", "text"):
                draw_path_pdf(c, el["commands"], el["color"], el.get("stroke"), el.get("width", 1.3))
            elif el["kind"] == "rect":
                c.setFillColor(el["color"])
                c.rect(el["x"], el["y"], el["w"], el["h"], stroke=0, fill=1)
            else:
                c.setStrokeColor(el["color"])
                c.setLineWidth(el["width"])
                c.line(el["x"], el["y"], el["x2"], el["y2"])
        c.restoreState()
        c.showPage()
    c.save()


def export_svg(slide: Slide, path: Path, width: int = W, height: int = H) -> None:
    elements = [f'<rect width="{width}" height="{height}" fill="{slide.bg}"/>']
    for el in slide.elements:
        if el["kind"] in ("path", "text"):
            d = " ".join(
                row[0] + " ".join(f"{v:.4f}" for v in row[1:]) for row in el["commands"]
            )
            label = html.escape(el.get("text", el.get("name", "Letterforms")), quote=True)
            elements.append(
                f'<path aria-label="{label}" d="{d}" fill="{el["color"] or "none"}"'
                f' stroke="{el.get("stroke") or "none"}" stroke-width="{el.get("width", 1.3)}"/>'
            )
        elif el["kind"] == "rect":
            elements.append(f'<rect x="{el["x"]}" y="{el["y"]}" width="{el["w"]}" height="{el["h"]}" fill="{el["color"]}"/>')
        else:
            elements.append(f'<path d="M{el["x"]},{el["y"]} L{el["x2"]},{el["y2"]}" stroke="{el["color"]}" stroke-width="{el["width"]}"/>')
    path.write_text(
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">'
        f"<title>{html.escape(slide.title)}</title>" + "".join(elements) + "</svg>",
        encoding="utf-8",
    )


def rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color.lstrip("#"))


def native_path(
    slide: Any, commands: list[Any], color: str | None, name: str,
    stroke: str | None = None, width: float = 1.3,
) -> None:
    """Create a native, editable PowerPoint freeform with exact Bézier curves."""
    if not any(row[0] != "Z" for row in commands):
        return
    xs = [row[i] for row in commands for i in range(1, len(row), 2)]
    ys = [row[i] for row in commands for i in range(2, len(row), 2)]
    x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
    unit = 1000
    builder = slide.shapes.build_freeform(0, 0, scale=SCALE / unit)
    builder.add_line_segments([(max(1, round((x1 - x0) * unit)), max(1, round((y1 - y0) * unit)))])
    item = builder.convert_to_shape(round(x0 * SCALE), round(y0 * SCALE))
    item.name = name[:150]
    if color:
        item.fill.solid()
        item.fill.fore_color.rgb = rgb(color)
    else:
        item.fill.background()
    if stroke:
        item.line.color.rgb = rgb(stroke)
        item.line.width = Pt(width * .6)
    else:
        item.line.fill.background()
    # PowerPoint's default freeform theme adds a drop shadow unless removed.
    style = item._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        item._element.remove(style)
    geometry = item._element.spPr.custGeom
    path_list = geometry.find("{http://schemas.openxmlformats.org/drawingml/2006/main}pathLst")
    assert path_list is not None
    for child in list(path_list):
        path_list.remove(child)
    dml_path = OxmlElement("a:path")
    dml_path.set("w", str(max(1, round((x1 - x0) * unit))))
    dml_path.set("h", str(max(1, round((y1 - y0) * unit))))
    dml_path.set("stroke", "1" if stroke else "0")
    if not color:
        dml_path.set("fill", "none")
    for row in commands:
        tag = {"M": "a:moveTo", "L": "a:lnTo", "C": "a:cubicBezTo", "Z": "a:close"}[row[0]]
        command = OxmlElement(tag)
        for i in range(1, len(row), 2):
            point = OxmlElement("a:pt")
            point.set("x", str(round((row[i] - x0) * unit)))
            point.set("y", str(round((row[i + 1] - y0) * unit)))
            command.append(point)
        dml_path.append(command)
    path_list.append(dml_path)


def export_pptx(slides: list[Slide], path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13 + 1 / 3), Inches(7.5)
    prs.core_properties.title = "New Orkhon — Typeface Presentation"
    prs.core_properties.subject = f"Experimental Old Turkic serif, release {VERSION}"
    prs.core_properties.author = "New Orkhon"
    for source in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(source.bg)
        slide.notes_slide.notes_text_frame.text = source.notes
        for el in source.elements:
            if el["kind"] == "path":
                native_path(
                    slide, el["commands"], el["color"], el["name"],
                    el.get("stroke"), el.get("width", 1.3),
                )
            elif el["kind"] == "text":
                # Instrument Sans's typographic ascent is 970 / 1000 em.
                ascent = .970 if Path(el["font"]) != MONO else 1.020
                box = slide.shapes.add_textbox(
                    round(el["x"] * SCALE), round((el["y"] - el["size"] * ascent) * SCALE),
                    round((el["w"] + el["size"] * .4) * SCALE), round(el["size"] * 1.5 * SCALE),
                )
                box.name = el["text"]
                frame = box.text_frame
                frame.margin_top = frame.margin_bottom = frame.margin_left = frame.margin_right = 0
                frame.vertical_anchor = MSO_ANCHOR.TOP
                frame.word_wrap = False
                paragraph = frame.paragraphs[0]
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.0
                run = paragraph.add_run()
                run.text = el["text"]
                tt, _ = font_data(el["font"])
                run.font.name = tt["name"].getDebugName(1)
                run.font.size = Pt(el["size"] * .6)
                run.font.color.rgb = rgb(el["color"])
                run._r.get_or_add_rPr().set("kern", "400000")
            elif el["kind"] == "rect":
                box = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, round(el["x"] * SCALE), round(el["y"] * SCALE),
                    round(el["w"] * SCALE), round(el["h"] * SCALE),
                )
                box.fill.solid()
                box.fill.fore_color.rgb = rgb(el["color"])
                box.line.fill.background()
                style = box._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
                if style is not None:
                    box._element.remove(style)
            else:
                native_path(
                    slide, [("M", el["x"], el["y"]), ("L", el["x2"], el["y2"])],
                    None, "Rule", el["color"], el["width"],
                )
    prs.save(str(path))


def case_ranges(words: list[str], capitals: set[int]) -> dict[str, Any]:
    """Apply lower case to everything but editorial capitals in logical RTL order."""
    ranges, offset = [], 0
    for number, word in enumerate(words):
        start = offset + (1 if number in capitals else 0)
        if start < offset + len(word):
            ranges.append((start, offset + len(word), 1))
        offset += len(word) + 1
    return {"ss01": ranges}


def wrapped_words(
    words: list[str], font: Path, size: float, max_width: float, max_lines: int,
    capitals: set[int] | None = None,
) -> list[tuple[int, list[str]]]:
    lines: list[tuple[int, list[str]]] = []
    current: list[str] = []
    start = 0
    for number, word in enumerate(words):
        proposal = " ".join([*current, word])
        features = None if capitals is None else case_ranges(
            [*current, word], {n - start for n in capitals if start <= n <= number},
        )
        width = shape(font, proposal, features)[1] * size / font_data(str(font))[0]["head"].unitsPerEm
        if width > max_width and current:
            lines.append((start, current))
            start, current = number, [word]
        else:
            current.append(word)
    if current:
        lines.append((start, current))
    assert len(lines) <= max_lines, f"{len(lines)} text lines exceeds {max_lines}"
    line_count = len(lines)

    @lru_cache
    def span_width(start: int, end: int) -> float:
        group = words[start:end]
        features = None if capitals is None else case_ranges(
            group, {n - start for n in capitals if start <= n < end},
        )
        return shape(font, " ".join(group), features)[1] * size / font_data(str(font))[0]["head"].unitsPerEm

    @lru_cache
    def balanced(start: int, remaining: int) -> tuple[float, tuple[int, ...]]:
        if remaining == 0:
            return (0., ()) if start == len(words) else (float("inf"), ())
        best: tuple[float, tuple[int, ...]] = (float("inf"), ())
        for end in range(start + 1, len(words) - remaining + 2):
            width = span_width(start, end)
            if width > max_width:
                break
            tail_cost, tail = balanced(end, remaining - 1)
            cost = (max_width - width) ** 2 + tail_cost
            if end - start == 1:
                cost += max_width ** 2
            if cost < best[0]:
                best = (cost, (end, *tail))
        return best

    cost, ends = balanced(0, line_count)
    assert cost < float("inf") and len(ends) == line_count
    result, start = [], 0
    for end in ends:
        result.append((start, words[start:end]))
        start = end
    return result


def dimension(
    slide: Slide, x1: float, y1: float, x2: float, y2: float, label: str,
) -> None:
    slide.line(x1, y1, x2, y2, width=1)
    vertical = abs(x2 - x1) < abs(y2 - y1)
    for x, y in ((x1, y1), (x2, y2)):
        slide.line(x - 6 if vertical else x, y if vertical else y - 6,
                   x + 6 if vertical else x, y if vertical else y + 6, width=1)
    if vertical:
        slide.text(label, x1 - 14, (y1 + y2) / 2 + 7, 21, font=MONO, align="right")
    else:
        slide.text(label, (x1 + x2) / 2, y1 - 14, 21, font=MONO, align="center")


def compose() -> list[Slide]:
    from glyph_similarity import outline as compiled_outline
    from shapely.geometry import LineString

    slides: list[Slide] = []
    passage = historical("D1")
    words, capitals = passage["words"], set(passage["caps"])
    s = Slide("New Orkhon / Kül Tigin D1")
    s.furniture(1, f"Text specimen / {VERSION}")
    s.text("New Orkhon", 82, 181, 108)
    s.text("Orkhon script reimagined for a modern era.", 88, 238, 28)
    s.line(88, 274, 1512, 274)
    breaks = wrapped_words(words, UPPER, 75, 1404, 4, capitals)
    for row, (start, group) in enumerate(breaks):
        caps = {n - start for n in capitals if start <= n < start + len(group)}
        s.run(" ".join(group), 1512, 382 + row * 105, 75, INK, UPPER,
              case_ranges(group, caps), "right", f"Full Kül Tigin D1 / {row + 1}")
    s.text("KÜL TİGİN / EAST FACE / D1", 88, 794, 17, font=MONO)
    s.text("Editorial case · typeset line breaks", 1512, 794, 19, align="right")
    s.notes = (
        "Complete Kül Tigin east-face D1 passage, 31 tokens. Source: "
        "https://tr.wikisource.org/wiki/Orhun_Yaz%C4%B1tlar%C4%B1_(K%C3%BCl_Tigin). "
        "The showcase source documents the Yenisei-to-Orkhon EZ substitution. "
        "Dividers are normalized to spaces. Capitals and line breaks are editorial. "
        "New Orkhon is speculative type design, not a recovered historical stage."
    )
    slides.append(s)

    s = Slide("New Orkhon Sans / Kül Tigin D1")
    s.furniture(2, f"Sans / {VERSION}")
    s.text("New Orkhon Sans", 82, 181, 108)
    s.text("Kül Tigin / east face / D1", 88, 238, 28)
    s.line(88, 274, 1512, 274)
    breaks = wrapped_words(words, SANS_UPPER, 75, 1404, 4, capitals)
    for row, (start, group) in enumerate(breaks):
        caps = {n - start for n in capitals if start <= n < start + len(group)}
        s.run(" ".join(group), 1512, 382 + row * 105, 75, INK, SANS_UPPER,
              case_ranges(group, caps), "right", f"Full Kül Tigin D1 / Sans / {row + 1}")
    s.text("42 SIGNS / TWO CASES / LOGICAL RTL", 88, 794, 17, font=MONO)
    s.text("Editorial case · typeset line breaks", 1512, 794, 19, align="right")
    s.notes = (
        "The complete passage from slide 1, now in the separate New Orkhon Sans family. "
        "Ordinary caps use 84/76 construction, lowercase 76/70 with dense-counter exceptions. "
        "Original outlines: no Helvetica glyph data is used."
    )
    slides.append(s)

    s = Slide("Serif / Sans", INK)
    s.furniture(3, "The two families", True)
    s.text("Serif / Sans", 82, 204, 104, PAPER)
    s.line(88, 266, 1512, 266, PAPER)
    s.line(800, 300, 800, 755, PAPER, .7)
    features = case_ranges([words[2]], {0})
    maximum = max(shape(font, words[2], features)[1] for font in (UPPER, SANS_UPPER))
    common_size = min(330, 610000 / maximum)
    for center, label, path in [(408, "SERIF", UPPER), (1192, "SANS", SANS_UPPER)]:
        s.text(label, center - 320, 333, 18, PAPER, MONO)
        s.run(words[2], center, 638, common_size, PAPER, path, features, "center")
    s.text("Teŋri / Kül Tigin D1", 88, 793, 20, PAPER)
    s.notes = "Both specimens use the same Unicode, case ranges, baseline and em size."
    slides.append(s)

    s = Slide("Kül Tigin / text comparison")
    s.furniture(4, "Text detail")
    s.text("Kül Tigin / D1", 82, 204, 94)
    s.text("Üze kök Teŋri", 88, 260, 25)
    excerpt = " ".join(words[:3])
    s.text("NOTO SANS OLD TURKIC / DIGITAL REFERENCE", 88, 309, 16, font=MONO)
    s.run(excerpt, 1512, 475, 212, INK, REFERENCE, align="right")
    s.line(88, 525, 1512, 525)
    s.text("NEW ORKHON / EDITORIAL CASE", 88, 575, 16, font=MONO)
    s.run(excerpt, 1512, 742, 212, INK, UPPER, case_ranges(words[:3], {0, 2}), "right")
    s.notes = (
        "Both rows use the same logical Unicode sequence. Noto Sans Old Turkic "
        "is a digital reference, not a facsimile of the inscription. New Orkhon "
        "adds an authored case distinction without changing the codepoints."
    )
    slides.append(s)

    s = Slide("Letter construction")
    s.furniture(5, "Serif construction")
    s.text("Letter construction", 82, 204, 94)
    for i, (code, name, caption) in enumerate([
        (0x10C24, "AEN", "Two uprights and one diagonal"),
        (0x10C22, "EM", "Open bowl and diagonal leg"),
        (0x10C43, "AT", "Side bowl and overhanging top arm"),
    ]):
        x = 88 + 496 * i
        s.text(f"{name} / U+{code:05X}", x, 297, 18, font=MONO)
        s.line(x, 319, x + 432, 319)
        s.text("REFERENCE", x, 358, 14, font=MONO)
        s.text("NEW ORKHON", x + 260, 358, 14, font=MONO)
        s.run(chr(code), x + 92, 589, 220, INK, REFERENCE, align="center", outline=True)
        s.run(chr(code), x + 338, 620, 310, INK, UPPER, align="center")
        s.text(caption, x, 731, 22)
    s.notes = (
        "Reference outlines are from Noto Sans Old Turkic. New forms are authored "
        "design choices. AEN uses a single N; EM follows the selected R-like writing "
        "route without a left post; AT has a b-like stem/bowl and a left-overhanging "
        "top arm. No historical inevitability is claimed."
    )
    slides.append(s)

    s = Slide("Capitals and lowercase")
    s.furniture(6, "Serif cases")
    s.text("Capitals and lowercase", 82, 204, 88)
    s.text("42 pairs · OpenType case alternates", 88, 262, 25)
    s.line(88, 620, 1512, 620, width=.7)
    for i, (code, name, caption) in enumerate([
        (0x10C24, "AEN", "Single N / n"),
        (0x10C0D, "AG", "Separate horns / cursive zhe"),
        (0x10C1E, "AL", "Stem / dotless j"),
        (0x10C32, "EC", "Capital lambda / curved head"),
    ]):
        x = 88 + 356 * i
        s.text(name, x, 349, 17, font=MONO)
        s.run(chr(code), x + 79, 620, 245, INK, UPPER, align="center", outline=True)
        s.run(chr(code), x + 250, 620, 245, INK, LOWER, align="center")
        s.text(caption, x, 736, 20)
        s.text("CAPITAL / LOWERCASE", x, 780, 14, font=MONO)
    s.notes = (
        "Case is selected with ss01/ss02 or separate default fonts, not invented "
        "Unicode case points. In 0.13 capital EC uses Greek capital-Lambda construction. "
        "In 0.13 its curved lowercase body moved up exactly 190 units. The 0.14 "
        "revision refits the serifs to its legs. Both feet stop on the baseline."
    )
    slides.append(s)

    s = Slide("Vertical metrics / blueprint 01")
    s.furniture(7, "Blueprint 01 / Serif")
    s.text("Vertical metrics", 82, 193, 90)
    s.text("Font units · 1000 units per em", 88, 246, 23)
    baseline, size = 650, 520
    for value, label in [(700, "700 / CAP"), (510, "510 / BODY"), (0, "0 / BASELINE"), (-190, "-190 / DESCENDER")]:
        y = baseline - value * size / 1000
        s.line(280, y, 1512, y, width=.6)
        s.text(label, 88, y - 10, 16, font=MONO)
    for code, font, center, name in [
        (0x10C00, UPPER, 472, "A / CAPITAL"),
        (0x10C32, UPPER, 918, "EC / CAPITAL"),
        (0x10C32, LOWER, 1305, "EC / LOWERCASE"),
    ]:
        _, bounds = s.run(chr(code), center, baseline, size, INK, font,
                          align="center", outline=True, width=1.6)
        s.line(bounds[0], 272, bounds[0], 771, width=.55)
        s.line(bounds[2], 272, bounds[2], 771, width=.55)
        s.text(name, center, 801, 15, font=MONO, align="center")
    s.notes = (
        "Contours are from the compiled fonts. Vertical hairlines mark ink bounds, "
        "not advance widths. Round forms may overshoot by 10 units. EC lowercase "
        "is an ascender-height exception: its head is at 700 and its feet at 0."
    )
    slides.append(s)

    for font_path, page, plate, style, nominal_stem, nominal_hair in [
        (UPPER, 8, 2, "Serif", 90, 64), (SANS_UPPER, 9, 3, "Sans", 84, 76),
    ]:
        s = Slide(f"{style} stroke weights / blueprint {plate:02d}")
        s.furniture(page, f"Blueprint {plate:02d} / {style}")
        s.text(f"{style} stroke weights", 82, 193, 90)
        s.text("Capital construction · font units", 88, 246, 23)
        s.text(str(nominal_stem), 88, 465, 120)
        s.text("Stem", 88, 510, 25)
        s.text(str(nominal_hair), 88, 661, 120)
        s.text("Thin stroke", 88, 706, 25)
        baseline, factor = 690, .54
        for code, x in [(0x10C03, 560), (0x10C00, 1060)]:
            s.run(chr(code), x, baseline, 540, INK, font_path, outline=True, width=1.7)
        tt, _ = font_data(str(font_path))
        stem = compiled_outline(tt, "u10C03").intersection(LineString([(-1000, 350), (2000, 350)]))
        x0, _, x1, _ = stem.bounds
        for x in (560 + x0 * factor, 560 + x1 * factor):
            s.line(x, baseline - 350 * factor, x, 574, width=.7)
        dimension(s, 560 + x0 * factor, 574, 560 + x1 * factor, 574, str(round(x1 - x0)))
        a = compiled_outline(tt, "u10C00")
        bars = a.intersection(LineString([(250, -100), (250, 800)]))
        top_bar = max(bars.geoms, key=lambda part: part.bounds[3])
        _, y0, _, y1 = top_bar.bounds
        for y in (baseline - y0 * factor, baseline - y1 * factor):
            s.line(1060 + 250 * factor, y, 1450, y, width=.7)
        dimension(s, 1450, baseline - y1 * factor, 1450, baseline - y0 * factor, str(round(y1 - y0)))
        s.text("I / U+10C03", 560, 775, 16, font=MONO)
        s.text("A / U+10C00", 1060, 775, 16, font=MONO)
        s.notes = (
            "Measurements are from compiled contours: I at y=350 and A at x=250. "
            f"{nominal_stem}/{nominal_hair} are ordinary {style} construction widths. "
            "Curved and diagonal sections have optical exceptions."
        )
        slides.append(s)

    s = Slide("AL kerning / blueprint 04")
    s.furniture(10, "Blueprint 04 / Serif")
    s.text("AL kerning", 82, 193, 90)
    s.text("The left neighbor can move over the hook when its underside is clear.", 88, 252, 24)
    for center, feature, label in [(408, {"kern": 0}, "KERN OFF"), (1192, None, "KERN ON")]:
        x = center - 320
        s.text(label, x, 321, 17, font=MONO)
        for y in (640, 640 - .41 * 510):
            s.line(x, y, x + 640, y, width=.6)
        s.run("𐰞𐰖", center, 640, 410, INK, LOWER, feature, "center", outline=True, width=1.6)
    off, on = shape(LOWER, "𐰞𐰖", {"kern": 0})[1], shape(LOWER, "𐰞𐰖")[1]
    s.text("AL + AY", 88, 791, 17, font=MONO)
    s.text(f"{round(off - on)} units less advance", 1512, 791, 24, align="right")
    s.notes = (
        "Logical AL + AY is physically AY to the left of AL because the script is RTL. "
        "73 safe case-specific neighbors are enabled in 0.13. The newly raised "
        "lowercase EC now also qualifies. Compiled clearance across all enabled pairs "
        "is at least 38.5 units. No standalone shape is changed by kerning."
    )
    slides.append(s)

    s = Slide("EM + I / blueprint 05")
    s.furniture(11, "Blueprint 05 / Serif")
    s.text("EM + I", 82, 193, 90)
    s.text("U+10C22 + U+10C03 · OpenType liga", 88, 252, 23)
    for center, feature, label, outlined in [
        (408, {"liga": 0}, "SEPARATE", True), (1192, None, "LIGATURE", False),
    ]:
        s.text(label, center - 320, 321, 17, font=MONO)
        s.line(center - 320, 645, center + 320, 645, width=.6)
        s.run("𐰢𐰃", center, 645, 490, INK, LOWER, feature, "center", outline=outlined, width=1.7)
    s.text("The text remains two Unicode characters.", 88, 791, 25)
    s.notes = (
        "Standard liga substitutes matching capital or lowercase pairs. Mixed cases "
        "remain separate. Source text stays U+10C22 U+10C03. The slide artwork is "
        "converted to vector paths, not live copyable text."
    )
    slides.append(s)

    for page in range(2):
        s = Slide(f"Alphabet / {page + 1}")
        s.furniture(12 + page, "Serif alphabet")
        s.text("Serif alphabet", 82, 192, 90)
        s.text("42 encoded signs · capital / lowercase", 88, 245, 24)
        s.text(f"0{page + 1} / 02", 1512, 245, 17, font=MONO, align="right")
        for row in (0, 1):
            s.line(88, 443 + row * 164, 1512, 443 + row * 164, width=.7)
        for index, letter in enumerate(LETTERS[page * 21:(page + 1) * 21]):
            col, row = index % 7, index // 7
            x, baseline = 88 + col * 204, 410 + row * 164
            name = unicodedata.name(chr(letter.code)).split("ORKHON ", 1)[1]
            s.text(name, x + 4, baseline - 105, 16, font=MONO)
            s.run(chr(letter.code), x + 62, baseline, 135, INK, UPPER, align="center")
            s.run(chr(letter.code), x + 160, baseline, 135, INK, LOWER, align="center")
        s.notes = (
            f"Alphabet page {page + 1} of 2. Every encoded Orkhon sign in formal Unicode "
            "order, with capital on the left and lowercase on the right. Case forms "
            "are alternate glyphs under the same Unicode characters."
        )
        slides.append(s)

    s = Slide(f"New Orkhon / {VERSION}", INK)
    s.furniture(14, f"Release {VERSION}", True)
    s.text("New Orkhon", 82, 211, 112, PAPER)
    s.text("Yeni Orhun", 1512, 210, 26, PAPER, REGULAR, "right")
    s.line(88, 245, 1512, 245, PAPER)
    s.run(words[2], 800, 620, 385, PAPER, SANS_UPPER,
          case_ranges([words[2]], {0}), "center", "Teŋri / Kül Tigin D1")
    s.line(88, 756, 1512, 756, PAPER)
    for x, label in [(88, "SERIF / SANS"), (615, "42 SIGNS / TWO CASES"), (1120, "RTL / OPENTYPE / OFL")]:
        s.text(label, x, 804, 16, PAPER, MONO)
    s.notes = (
        f"New Orkhon {VERSION}. Four installable TTFs with ss01/ss02, kern and liga. "
        "Teŋri is from Kül Tigin D1 with editorial case. Typeface licensed under OFL 1.1. "
        "The supplied PowerPoint uses native editable vector glyphs and editable Latin text."
    )
    slides.append(s)
    assert len(slides) == TOTAL
    return slides


def render_previews(pdf: Path, stem: str = "slide") -> None:
    document = fitz.open(pdf)
    for i, page in enumerate(document, 1):
        page.get_pixmap(matrix=fitz.Matrix(1600 / page.rect.width, 1600 / page.rect.width),
                        alpha=False).save(PREVIEWS / f"{stem}-{i:02d}.png")
    columns, thumb_w, thumb_h, gap = 2, 640, 360, 24
    rows = (len(document) + columns - 1) // columns
    sheet = Image.new("RGB", (columns * thumb_w + gap * (columns + 1),
                             rows * (thumb_h + 35) + gap * (rows + 1)), PAPER)
    draw = ImageDraw.Draw(sheet)
    font = ImageFont.truetype(str(REGULAR), 16)
    for i in range(len(document)):
        x, y = gap + i % columns * (thumb_w + gap), gap + i // columns * (thumb_h + 35 + gap)
        image = Image.open(PREVIEWS / f"{stem}-{i + 1:02d}.png")
        image.thumbnail((thumb_w, thumb_h), Image.Resampling.LANCZOS)
        sheet.paste(image, (x, y))
        draw.text((x, y + thumb_h + 8), f"{i + 1:02d}", fill=INK, font=font)
    sheet.save(PREVIEWS / f"{stem}-contact-sheet.png")


def blueprint_catalog() -> list[Slide]:
    plates = []
    for page in range(14):
        style, cap_font, small_font = [
            ("Serif", UPPER, LOWER), ("Sans", SANS_UPPER, SANS_LOWER),
        ][page // 7]
        local_page = page % 7
        s = Slide(f"{style} outlines / {local_page + 1}")
        s.text(f"New Orkhon / {style} outlines", 88, 140, 65)
        s.text(f"RELEASE {VERSION} / CAPITAL AND LOWERCASE", 88, 191, 17, font=MONO)
        for index, letter in enumerate(LETTERS[local_page * 6:(local_page + 1) * 6]):
            col, row = index % 3, index // 3
            x, baseline = 88 + col * 492, 427 + row * 280
            name = unicodedata.name(chr(letter.code)).split("ORKHON ", 1)[1]
            s.text(f"{name} / U+{letter.code:05X}", x, 253 + row * 280, 19, font=MONO)
            for value in (700, 510, 0, -190):
                y = baseline - value * .17
                s.line(x, y, x + 430, y, width=.5)
            s.run(chr(letter.code), x + 100, baseline, 170, INK, cap_font,
                  align="center", outline=True, width=1.2)
            s.run(chr(letter.code), x + 330, baseline, 170, INK, small_font,
                  align="center", outline=True, width=1.2)
            s.text("CAPITAL / LOWERCASE", x, baseline + 66, 14, font=MONO)
        s.text("Guides: 700 / 510 / 0 / -190 font units", 88, 850, 16, font=MONO)
        s.text(f"{page + 1:02d} / 14", 1512, 850, 16, font=MONO, align="right")
        plates.append(s)
    return plates


def export_glyph_outlines() -> None:
    folder = OUT / "glyph-outlines"
    folder.mkdir(exist_ok=True)
    for old in folder.glob("u*.svg"):
        old.unlink()
    for style, cap_font, small_font in [
        ("serif", UPPER, LOWER), ("sans", SANS_UPPER, SANS_LOWER),
    ]:
        destination = folder / style
        destination.mkdir(exist_ok=True)
        for letter, font, suffix in [
            (letter, font, suffix)
            for letter in LETTERS for font, suffix in ((cap_font, ""), (small_font, ".lc"))
        ]:
            tt, _ = font_data(str(font))
            name = f"u{letter.code:05X}{suffix}"
            commands, _ = glyph_outline(str(font), name)
            advance = tt["hmtx"][name][0]
            d = " ".join(
                row[0] + " ".join(f"{v:.4f}" for v in row[1:])
                for row in transform(commands, 0, 0, 1, -1)
            )
            (destination / f"{name}.svg").write_text(
                f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="-60 -800 {advance + 120} 1100">'
                f'<title>{name} / New Orkhon {style} {VERSION} / advance {advance}</title>'
                f'<path d="{d}" fill="none" stroke="#000000" stroke-width="2"/></svg>',
                encoding="utf-8",
            )


def main() -> None:
    OUT.mkdir(exist_ok=True)
    PREVIEWS.mkdir(exist_ok=True)
    parser = argparse.ArgumentParser()
    parser.add_argument("--render", type=Path)
    args = parser.parse_args()
    if args.render:
        render_previews(args.render, "powerpoint")
        return
    slides = compose()
    for i, slide in enumerate(slides, 1):
        export_svg(slide, OUT / f"slide-{i:02d}.svg")
    export_pdf(slides, OUT / "New-Orkhon.pdf")
    export_pptx(slides, OUT / "New-Orkhon.pptx")
    scene = [asdict(slide) for slide in slides]
    for slide_data in scene:
        for element in slide_data["elements"]:
            if "font" in element:
                font_path = Path(element["font"])
                element["font"] = str(
                    Path("typeface") / font_path.relative_to(PROJECT)
                    if font_path.is_relative_to(PROJECT) else font_path.relative_to(ROOT)
                )
    (OUT / "slides.json").write_text(json.dumps(scene, ensure_ascii=False), encoding="utf-8")
    render_previews(OUT / "New-Orkhon.pdf")
    plates = blueprint_catalog()
    export_pdf(plates, OUT / "New-Orkhon-Outlines.pdf")
    with fitz.open(OUT / "New-Orkhon.pdf") as source, fitz.open() as target:
        target.insert_pdf(source, from_page=6, to_page=10)
        target.set_metadata({"title": "New Orkhon / Construction blueprints"})
        target.save(OUT / "New-Orkhon-Blueprints.pdf")
    for i, plate in enumerate(plates, 1):
        export_svg(plate, OUT / f"outline-plate-{i:02d}.svg")
    export_glyph_outlines()
    print(f"Built {len(slides)} slides, all vector. No raster assets.")


if __name__ == "__main__":
    main()
